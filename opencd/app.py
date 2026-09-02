import csv
import io
import sqlite3
from datetime import date, datetime

from flask import Flask, flash, g, redirect, render_template, request, send_file, url_for

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

app = Flask(__name__)
app.secret_key = "almoxarifado-secret-key"

DATABASE = "almoxarifado.db"


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DATABASE)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DATABASE)
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS categorias (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL UNIQUE
        );

        CREATE TABLE IF NOT EXISTS produtos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            descricao TEXT,
            categoria_id INTEGER REFERENCES categorias(id),
            fornecedor TEXT,
            localizacao TEXT,
            quantidade INTEGER NOT NULL DEFAULT 0,
            quantidade_minima INTEGER NOT NULL DEFAULT 0,
            preco REAL NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS movimentacoes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            produto_id INTEGER NOT NULL REFERENCES produtos(id),
            tipo TEXT NOT NULL CHECK (tipo IN ('entrada', 'saida')),
            quantidade INTEGER NOT NULL,
            data TEXT NOT NULL,
            observacao TEXT
        );
        """
    )
    db.commit()

    if db.execute("SELECT COUNT(*) FROM categorias").fetchone()[0] == 0:
        categorias = ["Mouse", "Teclado", "Monitor", "Cabo", "Fonte", "Armazenamento"]
        for nome in categorias:
            db.execute("INSERT INTO categorias (nome) VALUES (?)", (nome,))

    db.commit()
    db.close()


@app.route("/")
def dashboard():
    db = get_db()
    total_produtos = db.execute("SELECT COUNT(*) FROM produtos").fetchone()[0]
    total_itens = db.execute("SELECT COALESCE(SUM(quantidade), 0) FROM produtos").fetchone()[0]
    valor_estoque = db.execute(
        "SELECT COALESCE(SUM(quantidade * preco), 0) FROM produtos"
    ).fetchone()[0]
    estoque_baixo = db.execute(
        """
        SELECT p.*, c.nome AS categoria
        FROM produtos p LEFT JOIN categorias c ON c.id = p.categoria_id
        WHERE p.quantidade <= p.quantidade_minima
        ORDER BY p.quantidade ASC
        """
    ).fetchall()
    recentes = db.execute(
        """
        SELECT m.*, p.nome AS produto_nome
        FROM movimentacoes m JOIN produtos p ON p.id = m.produto_id
        ORDER BY m.id DESC LIMIT 10
        """
    ).fetchall()
    return render_template(
        "dashboard.html",
        total_produtos=total_produtos,
        total_itens=total_itens,
        valor_estoque=valor_estoque,
        estoque_baixo=estoque_baixo,
        recentes=recentes,
    )


@app.route("/produtos")
def produtos():
    db = get_db()
    busca = request.args.get("busca", "").strip()
    categoria = request.args.get("categoria", "")

    sql = """
        SELECT p.*, c.nome AS categoria
        FROM produtos p LEFT JOIN categorias c ON c.id = p.categoria_id
        WHERE 1=1
    """
    params = []
    if busca:
        sql += " AND (p.nome LIKE ? OR p.descricao LIKE ? OR p.fornecedor LIKE ?)"
        params += [f"%{busca}%"] * 3
    if categoria:
        sql += " AND p.categoria_id = ?"
        params.append(categoria)
    sql += " ORDER BY p.nome"

    lista = db.execute(sql, params).fetchall()
    categorias = db.execute("SELECT * FROM categorias ORDER BY nome").fetchall()
    return render_template("produtos.html", produtos=lista, categorias=categorias, busca=busca, categoria=categoria)


@app.route("/produtos/novo", methods=["GET", "POST"])
def produto_novo():
    return produto_form(None)


@app.route("/produtos/<int:produto_id>/editar", methods=["GET", "POST"])
def produto_editar(produto_id):
    return produto_form(produto_id)


def produto_form(produto_id):
    db = get_db()
    if produto_id is None:
        produto = None
    else:
        produto = db.execute("SELECT * FROM produtos WHERE id = ?", (produto_id,)).fetchone()
        if produto is None:
            flash("Produto não encontrado.", "erro")
            return redirect(url_for("produtos"))

    categorias = db.execute("SELECT * FROM categorias ORDER BY nome").fetchall()

    if request.method == "POST":
        nome = request.form.get("nome", "").strip()
        if not nome:
            flash("O nome do produto é obrigatório.", "erro")
        else:
            if produto_id is None:
                db.execute(
                    """
                    INSERT INTO produtos (nome, descricao, categoria_id, fornecedor, localizacao, quantidade, quantidade_minima, preco)
                    VALUES (?,?,?,?,?,?,?,?)
                    """,
                    (
                        nome,
                        request.form.get("descricao", "").strip(),
                        request.form.get("categoria_id") or None,
                        request.form.get("fornecedor", "").strip(),
                        request.form.get("localizacao", "").strip(),
                        int(request.form.get("quantidade") or 0),
                        int(request.form.get("quantidade_minima") or 0),
                        float(request.form.get("preco") or 0),
                    ),
                )
                flash("Produto cadastrado com sucesso!", "ok")
            else:
                db.execute(
                    """
                    UPDATE produtos
                    SET nome=?, descricao=?, categoria_id=?, fornecedor=?, localizacao=?, quantidade=?, quantidade_minima=?, preco=?
                    WHERE id=?
                    """,
                    (
                        nome,
                        request.form.get("descricao", "").strip(),
                        request.form.get("categoria_id") or None,
                        request.form.get("fornecedor", "").strip(),
                        request.form.get("localizacao", "").strip(),
                        int(request.form.get("quantidade") or 0),
                        int(request.form.get("quantidade_minima") or 0),
                        float(request.form.get("preco") or 0),
                        produto_id,
                    ),
                )
                flash("Produto atualizado com sucesso!", "ok")
            db.commit()
            return redirect(url_for("produtos"))

    return render_template("produto_form.html", produto=produto, categorias=categorias)


@app.post("/produtos/<int:produto_id>/excluir")
def produto_excluir(produto_id):
    db = get_db()
    db.execute("DELETE FROM movimentacoes WHERE produto_id = ?", (produto_id,))
    db.execute("DELETE FROM produtos WHERE id = ?", (produto_id,))
    db.commit()
    flash("Produto excluído.", "ok")
    return redirect(url_for("produtos"))


@app.route("/movimentacoes")
def movimentacoes():
    db = get_db()
    produto_id = request.args.get("produto", "")
    sql = """
        SELECT m.*, p.nome AS produto_nome
        FROM movimentacoes m JOIN produtos p ON p.id = m.produto_id
        WHERE 1=1
    """
    params = []
    if produto_id:
        sql += " AND m.produto_id = ?"
        params.append(produto_id)
    sql += " ORDER BY m.id DESC LIMIT 200"
    lista = db.execute(sql, params).fetchall()
    produtos = db.execute("SELECT id, nome FROM produtos ORDER BY nome").fetchall()
    return render_template("movimentacoes.html", movimentacoes=lista, produtos=produtos, produto_id=produto_id)


@app.route("/movimentacoes/nova", methods=["GET", "POST"])
def movimentacao_nova():
    db = get_db()
    if request.method == "POST":
        produto_id = request.form.get("produto_id")
        tipo = request.form.get("tipo")
        quantidade = int(request.form.get("quantidade") or 0)
        observacao = request.form.get("observacao", "").strip()

        produto = db.execute("SELECT * FROM produtos WHERE id = ?", (produto_id,)).fetchone()
        if produto is None:
            flash("Selecione um produto válido.", "erro")
        elif quantidade <= 0:
            flash("A quantidade deve ser maior que zero.", "erro")
        elif tipo == "saida" and produto["quantidade"] < quantidade:
            flash(
                f"Estoque insuficiente! Disponível: {produto['quantidade']} unidade(s).",
                "erro",
            )
        else:
            if tipo == "entrada":
                nova_quantidade = produto["quantidade"] + quantidade
            else:
                nova_quantidade = produto["quantidade"] - quantidade
            db.execute("UPDATE produtos SET quantidade = ? WHERE id = ?", (nova_quantidade, produto_id))
            db.execute(
                """
                INSERT INTO movimentacoes (produto_id, tipo, quantidade, data, observacao)
                VALUES (?,?,?,?,?)
                """,
                (
                    produto_id,
                    tipo,
                    quantidade,
                    datetime.now().strftime("%d/%m/%Y %H:%M"),
                    observacao,
                ),
            )
            db.commit()
            flash(f"{tipo.capitalize()} de {quantidade} unidade(s) registrada!", "ok")
            return redirect(url_for("movimentacoes"))

    produtos = db.execute("SELECT id, nome, quantidade FROM produtos ORDER BY nome").fetchall()
    return render_template("movimentacao_form.html", produtos=produtos)


@app.post("/movimentacoes/<int:mov_id>/excluir")
def movimentacao_excluir(mov_id):
    db = get_db()
    mov = db.execute("SELECT * FROM movimentacoes WHERE id = ?", (mov_id,)).fetchone()
    if mov is None:
        flash("Movimentação não encontrada.", "erro")
    else:
        db.execute("DELETE FROM movimentacoes WHERE id = ?", (mov_id,))
        db.commit()
        flash("Registro de movimentação excluído.", "ok")
    return redirect(url_for("movimentacoes"))


@app.route("/categorias", methods=["GET", "POST"])
def categorias():
    db = get_db()
    if request.method == "POST":
        nome = request.form.get("nome", "").strip()
        if not nome:
            flash("O nome da categoria é obrigatório.", "erro")
        else:
            try:
                db.execute("INSERT INTO categorias (nome) VALUES (?)", (nome,))
                db.commit()
                flash("Categoria criada!", "ok")
            except sqlite3.IntegrityError:
                flash("Categoria já existe.", "erro")
    lista = db.execute("SELECT * FROM categorias ORDER BY nome").fetchall()
    return render_template("categorias.html", categorias=lista)


@app.post("/categorias/<int:cat_id>/excluir")
def categoria_excluir(cat_id):
    db = get_db()
    db.execute("UPDATE produtos SET categoria_id = NULL WHERE categoria_id = ?", (cat_id,))
    db.execute("DELETE FROM categorias WHERE id = ?", (cat_id,))
    db.commit()
    flash("Categoria excluída.", "ok")
    return redirect(url_for("categorias"))


def _larguras(headers):
    padrao = {
        "Produto": 28, "Categoria": 16, "Fornecedor": 22, "Local": 14,
        "Qtd. Estoque": 12, "Qtd. Mínima": 11, "Preço Unit. (R$)": 14,
        "Valor (R$)": 12, "Status": 10, "Data": 16, "Tipo": 10,
        "Quantidade": 12, "Observação": 30,
    }
    return [padrao.get(h, 18) for h in headers]


def _dados_relatorio(tipo, de_d, ate_d, produto_id):
    db = get_db()
    if tipo in ("estoque", "baixo"):
        sql = """
            SELECT p.*, c.nome AS categoria
            FROM produtos p LEFT JOIN categorias c ON c.id = p.categoria_id
        """
        if tipo == "baixo":
            sql += " WHERE p.quantidade <= p.quantidade_minima"
        sql += " ORDER BY p.nome"
        linhas = []
        for p in db.execute(sql).fetchall():
            if p["quantidade"] == 0:
                status = "ZERADO"
            elif p["quantidade"] <= p["quantidade_minima"]:
                status = "BAIXO"
            else:
                status = "OK"
            linhas.append([
                p["nome"], p["categoria"] or "-", p["fornecedor"] or "-",
                p["localizacao"] or "-", p["quantidade"], p["quantidade_minima"],
                p["preco"], round(p["quantidade"] * p["preco"], 2), status,
            ])
        headers = ["Produto", "Categoria", "Fornecedor", "Local", "Qtd. Estoque",
                   "Qtd. Mínima", "Preço Unit. (R$)", "Valor (R$)", "Status"]
        rodape = ["TOTAL", "", "", "", sum(r[4] for r in linhas), "",
                  "", round(sum(r[7] for r in linhas), 2), ""]
        return headers, linhas, rodape

    sql = """
        SELECT m.*, p.nome AS produto_nome
        FROM movimentacoes m JOIN produtos p ON p.id = m.produto_id
        WHERE 1=1
    """
    params = []
    if produto_id:
        sql += " AND m.produto_id = ?"
        params.append(produto_id)
    sql += " ORDER BY m.id DESC LIMIT 2000"
    linhas = []
    for m in db.execute(sql, params).fetchall():
        try:
            md = datetime.strptime(m["data"], "%d/%m/%Y %H:%M").date()
        except ValueError:
            md = None
        if de_d and (md is None or md < de_d):
            continue
        if ate_d and (md is None or md > ate_d):
            continue
        linhas.append([m["data"], m["produto_nome"], m["tipo"], m["quantidade"], m["observacao"] or "-"])
    headers = ["Data", "Produto", "Tipo", "Quantidade", "Observação"]
    tot_ent = sum(r[3] for r in linhas if r[2] == "entrada")
    tot_sai = sum(r[3] for r in linhas if r[2] == "saida")
    rodape = ["TOTAL", f"Entradas: {tot_ent} unidades | Saídas: {tot_sai} unidades", "", "", ""]
    return headers, linhas, rodape


def _pdf_relatorio(nome_arquivo, titulo, subtitulo, headers, linhas, rodape):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title=nome_arquivo,
    )
    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle("titulo", parent=estilos["Title"],
                                   fontName="Helvetica-Bold", fontSize=15, spaceAfter=2)
    estilo_sub = ParagraphStyle("sub", parent=estilos["Normal"], fontSize=9,
                                textColor=colors.grey, spaceAfter=10)
    estilo_celula = ParagraphStyle("celula", parent=estilos["Normal"], fontSize=8)
    estilo_marca = ParagraphStyle("marca", parent=estilos["Normal"], fontSize=9)

    dados = [[Paragraph(f"<b>{h}</b>", estilo_celula) for h in headers]]
    for r in linhas:
        dados.append([Paragraph(str(c), estilo_celula) for c in r])
    if rodape:
        dados.append([Paragraph(f"<b>{c}</b>", estilo_celula) for c in rodape])

    ncol = len(headers)
    tb = Table(dados, colWidths=[doc.width / ncol] * ncol, repeatRows=1)
    estilo_tb = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]
    if rodape and linhas:
        ult = len(dados) - 1
        estilo_tb.append(("BACKGROUND", (0, ult), (-1, ult), colors.HexColor("#e8edf2")))
        estilo_tb.append(("FONTNAME", (0, ult), (-1, ult), "Helvetica-Bold"))
    tb.setStyle(TableStyle(estilo_tb))

    doc.build([
        Paragraph(titulo, estilo_titulo),
        Paragraph(subtitulo, estilo_sub),
        tb,
        Spacer(0, 4),
        Paragraph(f"Total de registros: <b>{len(linhas)}</b>", estilo_marca),
    ])
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name=nome_arquivo,
                     mimetype="application/pdf")


def _xlsx_relatorio(nome_arquivo, titulo, subtitulo, headers, linhas, rodape):
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ncol = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = Font(bold=True, size=14, color="1E3A5F")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws.cell(row=2, column=1, value=subtitulo).font = Font(italic=True, size=9, color="7F8C8D")

    linha = 4
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=linha, column=j, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(horizontal="center")
    for r in linhas:
        linha += 1
        for j, v in enumerate(r, start=1):
            cell = ws.cell(row=linha, column=j, value=v)
            if isinstance(v, (int, float)):
                cell.number_format = "R$ #,##0.00"
    if rodape:
        linha += 1
        for j, v in enumerate(rodape, start=1):
            cell = ws.cell(row=linha, column=j, value=v)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EDF2")
            if isinstance(v, (int, float)):
                cell.number_format = "R$ #,##0.00"

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(ncol)}{max(4, 4 + len(linhas))}"
    for j, w in enumerate(_larguras(headers), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _csv_relatorio(nome_arquivo, headers, linhas, rodape):
    buffer = io.StringIO()
    w = csv.writer(buffer, delimiter=";")
    w.writerow(headers)
    w.writerows(linhas)
    if rodape:
        w.writerow(rodape)
    dados = buffer.getvalue().encode("utf-8-sig")
    return send_file(io.BytesIO(dados), as_attachment=True, download_name=nome_arquivo,
                     mimetype="text/csv; charset=utf-8")


@app.route("/relatorios")
def relatorios():
    db = get_db()
    produtos = db.execute("SELECT id, nome FROM produtos ORDER BY nome").fetchall()
    return render_template(
        "relatorios.html", produtos=produtos,
        reportlab_ok=REPORTLAB_OK, openpyxl_ok=OPENPYXL_OK,
    )


@app.route("/relatorios/gerar")
def relatorios_gerar():
    tipo = request.args.get("tipo", "estoque")
    formato = request.args.get("formato", "pdf")
    de = request.args.get("de", "")
    ate = request.args.get("ate", "")
    produto = request.args.get("produto", "")

    def parse_data(s):
        try:
            return datetime.strptime(s, "%Y-%m-%d").date()
        except ValueError:
            return None

    if tipo not in ("estoque", "baixo", "movimentacoes"):
        flash("Tipo de relatório inválido.", "erro")
        return redirect(url_for("relatorios"))
    if formato not in ("pdf", "xlsx", "csv"):
        flash("Formato inválido.", "erro")
        return redirect(url_for("relatorios"))
    if formato == "pdf" and not REPORTLAB_OK:
        flash("Biblioteca reportlab não instalada. Rode: sudo apt install python3-reportlab", "erro")
        return redirect(url_for("relatorios"))
    if formato == "xlsx" and not OPENPYXL_OK:
        flash("Biblioteca openpyxl não instalada. Rode: sudo apt install python3-openpyxl", "erro")
        return redirect(url_for("relatorios"))

    headers, linhas, rodape = _dados_relatorio(tipo, parse_data(de), parse_data(ate), produto)

    titulos = {
        "estoque": "Inventário de Estoque",
        "baixo": "Estoque Baixo / Zerado",
        "movimentacoes": "Movimentações de Estoque",
    }
    periodo = ""
    if tipo == "movimentacoes":
        if de and ate:
            periodo = f" | Período: {de} a {ate}"
        elif de:
            periodo = f" | A partir de: {de}"
        elif ate:
            periodo = f" | Até: {ate}"
    subtitulo = f"Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}{periodo} | Registros: {len(linhas)}"

    base = {"estoque": "inventario", "baixo": "estoque_baixo", "movimentacoes": "movimentacoes"}[tipo]
    nome_arquivo = f"relatorio_{base}_{date.today().isoformat()}.{formato}"

    if formato == "pdf":
        return _pdf_relatorio(nome_arquivo, titulos[tipo], subtitulo, headers, linhas, rodape)
    if formato == "xlsx":
        return _xlsx_relatorio(nome_arquivo, titulos[tipo], subtitulo, headers, linhas, rodape)
    return _csv_relatorio(nome_arquivo, headers, linhas, rodape)


init_db()

if __name__ == "__main__":
    app.run(debug=True)
